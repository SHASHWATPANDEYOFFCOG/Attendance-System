"""Authority-side student management: directory, detail, approvals, edits, bulk import.

Every query is scoped with `student_scope_clause` (in SQL) and every write goes through
`assert_can_manage_student`, so a teacher cannot reach a student outside their classes by
guessing an id.
"""
from __future__ import annotations

import csv
import io
import json
import secrets
from datetime import timedelta
from typing import Dict, List, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app import analytics_service as analytics
from app import validation as V
from app.audit import audit
from app.auth import is_admin, is_super_admin, require_admin, require_staff
from app.clock import tz_of
from app.crypto import hash_password, token_digest
from app.db import get_db
from app.models import (
    AttendanceRecordRow,
    Batch,
    COUNTS_PRESENT,
    Department,
    EXCLUDED_FROM_TOTAL,
    LeaveRequest,
    PasswordResetToken,
    Program,
    ProfileEditRequest,
    ReviewStatus,
    Role,
    Section,
    StudentProfile,
    StudentStatus,
    User,
    utcnow,
)
from app.notify import notify
from app.routes.academic import section_label
from app.routes.accounts import (
    apply_profile_changes,
    check_unique,
    profile_out,
    resolve_class,
    _clean_self,
)
from app.runtime import cfg
from app.scope import (
    assert_can_manage_section,
    assert_can_manage_student,
    manageable_section_ids,
    student_scope_clause,
)
from app import settings_service
from app.uploads import MB, read_limited

router = APIRouter(prefix="/api", tags=["students"])

INVITE_TTL = timedelta(days=7)


def _threshold(db: Session) -> float:
    return float(settings_service.get(db, cfg, "min_attendance_pct"))


def _pct_map(db: Session, user_ids: List[int]) -> Dict[int, Optional[float]]:
    """Overall attendance % per student, computed in SQL with the shared status rules."""
    if not user_ids:
        return {}
    eff = func.coalesce(AttendanceRecordRow.overridden_status, AttendanceRecordRow.status)
    from sqlalchemy import case

    counted = func.sum(case((eff.in_(tuple(EXCLUDED_FROM_TOTAL)), 0), else_=1))
    present = func.sum(case((eff.in_(tuple(COUNTS_PRESENT)), 1), else_=0))
    rows = db.execute(
        select(AttendanceRecordRow.user_id, counted, present)
        .where(AttendanceRecordRow.user_id.in_(user_ids),
               AttendanceRecordRow.deleted_at.is_(None))
        .group_by(AttendanceRecordRow.user_id)).all()
    out: Dict[int, Optional[float]] = {uid: None for uid in user_ids}
    for uid, c, p in rows:
        out[uid] = round(100.0 * (p or 0) / c, 1) if c else None
    return out


def _row(u: User, prof: StudentProfile, pct: Optional[float]) -> dict:
    sec = prof.section
    return {
        "id": u.id, "full_name": u.display_name, "email": u.email, "phone": u.phone,
        "roll_number": prof.roll_number or u.roll_number, "enrollment_no": prof.enrollment_no,
        "department": prof.department.code if prof.department else
        (sec.program.department.code if sec else None),
        "batch": prof.batch.label if prof.batch else None,
        "year": prof.year or (sec.year if sec else None),
        "section_id": prof.section_id, "section": section_label(sec) if sec else None,
        "status": prof.status, "email_verified": u.email_verified_at is not None,
        "has_face": u.has_biometrics, "attendance_pct": pct,
        "photo_url": f"/api/files/{prof.photo_file_id}" if prof.photo_file_id else None,
        "created_at": u.created_at,
    }


# ------------------------------------------------------------------ directory

@router.get("/students")
def list_students(
    q: Optional[str] = None, department_id: Optional[int] = None,
    batch_id: Optional[int] = None, year: Optional[int] = None,
    section_id: Optional[int] = None, status: Optional[str] = None,
    min_pct: Optional[float] = None, max_pct: Optional[float] = None,
    page: int = 1, page_size: int = 50,
    db: Session = Depends(get_db), actor: User = Depends(require_staff),
):
    stmt = (select(User, StudentProfile)
            .join(StudentProfile, StudentProfile.user_id == User.id)
            .where(User.deleted_at.is_(None), student_scope_clause(db, actor)))
    if q:
        like = f"%{q.strip()}%"
        stmt = stmt.where(or_(User.display_name.ilike(like), User.roll_number.ilike(like),
                              StudentProfile.roll_number.ilike(like), User.email.ilike(like),
                              StudentProfile.enrollment_no.ilike(like)))
    if department_id:
        stmt = stmt.where(StudentProfile.department_id == department_id)
    if batch_id:
        stmt = stmt.where(StudentProfile.batch_id == batch_id)
    if year:
        stmt = stmt.where(StudentProfile.year == year)
    if section_id:
        stmt = stmt.where(StudentProfile.section_id == section_id)
    if status:
        stmt = stmt.where(StudentProfile.status == status)
    stmt = stmt.order_by(StudentProfile.roll_number, User.display_name)

    page_size = max(10, min(200, page_size))
    if min_pct is not None or max_pct is not None:
        # Percentage is derived, so filter after computing it for the scoped set.
        rows = list(db.execute(stmt).all())
        pcts = _pct_map(db, [u.id for u, _ in rows])
        lo = min_pct if min_pct is not None else -1
        hi = max_pct if max_pct is not None else 101
        rows = [(u, p) for u, p in rows if pcts[u.id] is not None and lo <= pcts[u.id] <= hi]
        total = len(rows)
        rows = rows[(page - 1) * page_size: page * page_size]
    else:
        total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
        rows = list(db.execute(stmt.offset((max(1, page) - 1) * page_size).limit(page_size)).all())
        pcts = _pct_map(db, [u.id for u, _ in rows])
    return {"rows": [_row(u, p, pcts.get(u.id)) for u, p in rows],
            "meta": {"page": max(1, page), "page_size": page_size, "total": total,
                     "pages": (total + page_size - 1) // page_size},
            "threshold_pct": _threshold(db)}


@router.get("/students/{student_id}")
def student_detail(student_id: int, db: Session = Depends(get_db),
                   actor: User = Depends(require_staff)):
    prof = assert_can_manage_student(db, actor, student_id)
    user = db.get(User, student_id)
    out = profile_out(db, user, prof)
    summary = analytics.student_summary(db, user, threshold=_threshold(db), tz=tz_of(cfg))
    summary["calendar"] = analytics.attendance_calendar(db, user.id, days=120)
    out["attendance"] = summary
    out["leaves"] = [
        {"id": l.id, "type": l.type, "from_date": l.from_date.isoformat(),
         "to_date": l.to_date.isoformat(), "reason": l.reason, "status": l.status,
         "comment": l.comment, "created_at": l.created_at}
        for l in db.scalars(select(LeaveRequest).where(LeaveRequest.requester_id == student_id)
                            .order_by(LeaveRequest.created_at.desc()).limit(50))]
    out["edit_requests"] = [
        {"id": r.id, "changes": json.loads(r.changes), "reason": r.reason, "status": r.status,
         "remark": r.remark, "created_at": r.created_at}
        for r in db.scalars(select(ProfileEditRequest).where(
            ProfileEditRequest.user_id == student_id)
            .order_by(ProfileEditRequest.created_at.desc()).limit(20))]
    return out


# ------------------------------------------------------------------ approvals

class DecisionIn(BaseModel):
    reason: Optional[str] = Field(None, max_length=300)


@router.post("/students/{student_id}/approve")
def approve_student(student_id: int, request: Request, db: Session = Depends(get_db),
                    actor: User = Depends(require_staff)):
    prof = assert_can_manage_student(db, actor, student_id)
    user = db.get(User, student_id)
    if prof.status == StudentStatus.APPROVED.value:
        raise HTTPException(409, "This student is already approved.")
    if user.email and user.email_verified_at is None:
        raise HTTPException(409, "The student has not verified their email address yet.")
    prof.status = StudentStatus.APPROVED.value
    prof.rejection_reason = None
    prof.decided_by_id, prof.decided_at = actor.id, utcnow()
    notify(db, [user.id], "registration_approved", "Registration approved",
           "Your account has been approved. You now have full access, including marking "
           "attendance.", link="/student", email=True)
    audit(db, actor, "student_approved", target_type="user", target_id=user.id, request=request)
    db.commit()
    return {"id": user.id, "status": prof.status}


@router.post("/students/{student_id}/reject")
def reject_student(student_id: int, body: DecisionIn, request: Request,
                   db: Session = Depends(get_db), actor: User = Depends(require_staff)):
    prof = assert_can_manage_student(db, actor, student_id)
    reason = (body.reason or "").strip()
    if len(reason) < 3:
        raise HTTPException(422, detail={"field": "reason",
                                         "message": "Give the student a reason for the rejection."})
    if prof.status == StudentStatus.REJECTED.value:
        raise HTTPException(409, "This registration is already rejected.")
    was_approved = prof.status == StudentStatus.APPROVED.value
    prof.status = StudentStatus.REJECTED.value
    prof.rejection_reason = reason
    prof.decided_by_id, prof.decided_at = actor.id, utcnow()
    notify(db, [student_id], "registration_rejected", "Registration rejected",
           f"Your registration was not approved. Reason: {reason}. You can correct your "
           f"details from your profile and resubmit.", link="/student/profile", email=True)
    audit(db, actor, "student_rejected", target_type="user", target_id=student_id,
          reason=reason, detail={"was_approved": was_approved}, request=request)
    db.commit()
    return {"id": student_id, "status": prof.status}


@router.get("/approvals/summary")
def approvals_summary(db: Session = Depends(get_db), actor: User = Depends(require_staff)):
    from app.models import Enrollment, FaceReenrollRequest

    scope = student_scope_clause(db, actor)
    base = select(func.count()).select_from(StudentProfile).join(
        User, User.id == StudentProfile.user_id).where(User.deleted_at.is_(None), scope)
    pending_students = db.scalar(base.where(StudentProfile.status == "pending",
                                            User.email_verified_at.is_not(None))) or 0
    edit_requests = db.scalar(
        select(func.count()).select_from(ProfileEditRequest)
        .join(StudentProfile, StudentProfile.user_id == ProfileEditRequest.user_id)
        .where(ProfileEditRequest.status == "pending", scope)) or 0
    faces = db.scalar(
        select(func.count()).select_from(Enrollment)
        .join(StudentProfile, StudentProfile.user_id == Enrollment.user_id)
        .where(Enrollment.status == "pending", Enrollment.is_active.is_(True), scope)) or 0
    reenroll = db.scalar(
        select(func.count()).select_from(FaceReenrollRequest)
        .join(StudentProfile, StudentProfile.user_id == FaceReenrollRequest.user_id)
        .where(FaceReenrollRequest.status == "pending", scope)) or 0
    staff = 0
    if is_admin(actor):
        from app.models import StaffRegistration
        from app.routes.staff_signup import scoped_registrations

        staff = db.scalar(select(func.count()).select_from(
            scoped_registrations(db, actor).join(User, User.id == StaffRegistration.user_id)
            .where(StaffRegistration.status == "pending", User.email_verified_at.is_not(None))
            .subquery())) or 0
    return {"registrations": pending_students, "edit_requests": edit_requests,
            "face_enrollments": faces, "reenroll_requests": reenroll, "staff": staff,
            "can_review_staff": is_admin(actor),
            "total": pending_students + edit_requests + faces + reenroll + staff}


# ------------------------------------------------------------ edit requests

@router.get("/edit-requests")
def list_edit_requests(status: str = "pending", db: Session = Depends(get_db),
                       actor: User = Depends(require_staff)):
    stmt = (select(ProfileEditRequest, User, StudentProfile)
            .join(User, User.id == ProfileEditRequest.user_id)
            .join(StudentProfile, StudentProfile.user_id == User.id)
            .where(student_scope_clause(db, actor), ProfileEditRequest.status == status)
            .order_by(ProfileEditRequest.created_at))
    out = []
    for r, u, p in db.execute(stmt).all():
        current = {"full_name": u.display_name, "roll_number": p.roll_number,
                   "enrollment_no": p.enrollment_no, "department_id": p.department_id,
                   "batch_id": p.batch_id, "year": p.year, "section_id": p.section_id}
        out.append({"id": r.id, "student_id": u.id, "student": u.display_name,
                    "roll_number": p.roll_number, "changes": json.loads(r.changes),
                    "current": current, "reason": r.reason, "status": r.status,
                    "created_at": r.created_at})
    return out


@router.post("/edit-requests/{request_id}/{decision}")
def decide_edit_request(request_id: int, decision: str, body: DecisionIn, request: Request,
                        db: Session = Depends(get_db), actor: User = Depends(require_staff)):
    if decision not in ("approve", "reject"):
        raise HTTPException(404, "unknown action")
    r = db.get(ProfileEditRequest, request_id)
    if r is None:
        raise HTTPException(404, "request not found")
    prof = assert_can_manage_student(db, actor, r.user_id)
    if r.status != ReviewStatus.PENDING.value:
        raise HTTPException(409, "This request was already decided.")
    user = db.get(User, r.user_id)
    remark = (body.reason or "").strip() or None
    if decision == "reject" and not remark:
        raise HTTPException(422, detail={"field": "reason", "message": "Give a reason."})
    if decision == "approve":
        diff = apply_profile_changes(db, user, prof, _clean_self(json.loads(r.changes)))
        r.status = ReviewStatus.APPROVED.value
        audit(db, actor, "edit_request_approved", target_type="profile_edit_request",
              target_id=r.id, detail=diff, reason=remark, request=request)
    else:
        r.status = ReviewStatus.REJECTED.value
        audit(db, actor, "edit_request_rejected", target_type="profile_edit_request",
              target_id=r.id, reason=remark, request=request)
    r.decided_by_id, r.decided_at, r.remark = actor.id, utcnow(), remark
    notify(db, [user.id], f"edit_request_{r.status}",
           f"Profile edit request {r.status}",
           f"Your request to change {', '.join(json.loads(r.changes))} was {r.status}."
           + (f" Remark: {remark}" if remark else ""), link="/student/profile", email=True)
    db.commit()
    return {"id": r.id, "status": r.status}


# --------------------------------------------------------------- staff edits

class StaffStudentIn(BaseModel):
    full_name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    roll_number: Optional[str] = None
    enrollment_no: Optional[str] = None
    department_id: Optional[int] = None
    batch_id: Optional[int] = None
    year: Optional[int] = None
    section_id: Optional[int] = None
    dob: Optional[str] = None
    gender: Optional[str] = None
    guardian_name: Optional[str] = None
    guardian_phone: Optional[str] = None
    guardian_email: Optional[str] = None
    address: Optional[str] = None


@router.patch("/students/{student_id}")
def staff_update_student(student_id: int, body: StaffStudentIn, request: Request,
                         db: Session = Depends(get_db), actor: User = Depends(require_staff)):
    prof = assert_can_manage_student(db, actor, student_id)
    user = db.get(User, student_id)
    changes = body.model_dump(exclude_unset=True)
    email = changes.pop("email", None)
    if "section_id" in changes:
        assert_can_manage_section(db, actor, int(changes["section_id"]))
    clean = _clean_self(changes)
    diff = apply_profile_changes(db, user, prof, clean)
    if email is not None:
        try:
            new_email = V.email(email)
        except V.Invalid as exc:
            raise HTTPException(422, detail={"field": exc.field, "message": exc.message})
        if new_email != user.email:
            check_unique(db, email=new_email, exclude_user=user.id)
            diff["email"] = [user.email, new_email]
            user.email, user.email_verified_at = new_email, None
    if diff:
        audit(db, actor, "student_updated", target_type="user", target_id=user.id,
              detail=diff, request=request)
        notify(db, [user.id], "profile_changed", "Your profile was updated",
               f"{actor.display_name} updated: {', '.join(sorted(diff))}.",
               link="/student/profile")
    db.commit()
    return profile_out(db, user, prof)


@router.delete("/students/{student_id}")
def remove_student(student_id: int, request: Request, reason: str = Query(..., min_length=3),
                   db: Session = Depends(get_db), actor: User = Depends(require_admin)):
    """Graduated / left: the account is closed and biometrics destroyed; attendance history
    (an academic record) is kept."""
    from app import attendance_service as svc

    assert_can_manage_student(db, actor, student_id)
    user = db.get(User, student_id)
    svc.erase_biometrics(db, user, cfg, actor)
    user.is_active = False
    user.deleted_at = utcnow()
    audit(db, actor, "student_removed", target_type="user", target_id=student_id,
          reason=reason, detail={"biometrics_erased": True, "history_retained": True},
          request=request)
    db.commit()
    return {"removed": student_id}


# ------------------------------------------------------------- create / import

IMPORT_COLUMNS = ["full_name", "email", "mobile", "roll_number", "enrollment_no", "branch_code",
                  "batch", "year", "section", "dob", "gender", "guardian_name",
                  "guardian_phone", "guardian_email", "address"]
TEMPLATE_EXAMPLE = ["Asha Verma", "asha.verma@college.edu", "9876543210", "CSE23014",
                    "EN2023CSE014", "CSE", "2023-2027", "2", "A", "2005-04-17", "female",
                    "Rajesh Verma", "9812345678", "rajesh.verma@example.com",
                    "12 MG Road, Pune"]


@router.get("/students/import/template")
def import_template(_: User = Depends(require_staff)):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(IMPORT_COLUMNS)
    w.writerow(TEMPLATE_EXAMPLE)
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": 'attachment; filename="students-template.csv"'})


def _read_rows(data: bytes, filename: str) -> List[Dict[str, str]]:
    name = (filename or "").lower()
    if data[:4] == b"PK\x03\x04" or name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(422, "That Excel file could not be read. Save it as .xlsx.")
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        try:
            header = [str(h or "").strip().lower() for h in next(it)]
        except StopIteration:
            raise HTTPException(422, "The sheet is empty.")
        rows = []
        for values in it:
            if values is None or all(v in (None, "") for v in values):
                continue
            row = {}
            for h, v in zip(header, values):
                if hasattr(v, "isoformat"):
                    v = v.isoformat()[:10]
                elif isinstance(v, float) and v.is_integer():
                    v = int(v)
                row[h] = "" if v is None else str(v).strip()
            rows.append(row)
        return rows
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")   # Excel's default "CSV" on Windows
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise HTTPException(422, "The file has no header row.")
    return [{(k or "").strip().lower(): (v or "").strip() for k, v in r.items()}
            for r in reader if any((v or "").strip() for v in r.values())]


def process_rows(db: Session, actor: User, rows: List[Dict[str, str]], *, dry_run: bool,
                 send_invites: bool, source: str, request: Optional[Request]) -> dict:
    """Validate every row, then (unless dry_run) create the valid ones in one transaction.

    Shared by the bulk import and the single "add student" form, so the two can never apply
    different rules.
    """
    missing = [c for c in IMPORT_COLUMNS if c not in rows[0] and c != "guardian_email"]
    if missing:
        raise HTTPException(422, f"Missing column(s): {', '.join(missing)}. Download the template.")

    depts = {d.code.upper(): d for d in db.scalars(select(Department))}
    batches = {b.label: b for b in db.scalars(select(Batch))}
    allowed_sections = manageable_section_ids(db, actor)
    seen_rolls, seen_emails, seen_phones, seen_enrol = set(), set(), set(), set()
    errors: List[dict] = []
    valid: List[dict] = []

    for i, r in enumerate(rows, start=2):            # row 1 is the header
        def err(field: str, message: str):
            errors.append({"row": i, "field": field, "value": r.get(field, ""),
                           "message": message})

        try:
            rec = {
                "full_name": V.name(r.get("full_name")),
                "email": V.email(r.get("email")),
                "phone": V.phone(r.get("mobile")),
                "roll_number": V.roll(r.get("roll_number")),
                "enrollment_no": V.text(r.get("enrollment_no"), "enrollment_no", 48,
                                        required=True),
                "dob": V.dob(r.get("dob"), required=False),
                "gender": V.gender(r.get("gender"), required=False),
                "guardian_name": V.name(r.get("guardian_name"), "guardian_name", required=False),
                "guardian_phone": V.phone(r.get("guardian_phone"), "guardian_phone",
                                          required=False),
                "guardian_email": V.email(r.get("guardian_email"), "guardian_email",
                                          required=False),
                "address": V.text(r.get("address"), "address", 500),
            }
        except V.Invalid as exc:
            err(exc.field, exc.message)
            continue
        dept = depts.get((r.get("branch_code") or "").strip().upper())
        if dept is None:
            err("branch_code", f"Unknown branch {r.get('branch_code')!r}.")
            continue
        batch = batches.get((r.get("batch") or "").strip())
        if batch is None:
            err("batch", f"Unknown batch {r.get('batch')!r} (use e.g. 2023-2027).")
            continue
        try:
            year = int(r.get("year") or "")
        except ValueError:
            err("year", "Year must be a number 1-6.")
            continue
        sec = db.scalars(select(Section).join(Program, Program.id == Section.program_id).where(
            Program.department_id == dept.id, Section.year == year,
            func.upper(Section.name) == (r.get("section") or "").strip().upper(),
            or_(Section.batch_id == batch.id, Section.batch_id.is_(None)))).first()
        if sec is None:
            err("section", f"No section {r.get('section')!r} in {dept.code} year {year} "
                           f"batch {batch.label}.")
            continue
        if allowed_sections is not None and sec.id not in allowed_sections:
            err("section", f"{section_label(sec)} is not one of your classes.")
            continue
        if rec["enrollment_no"] in seen_enrol or db.scalars(select(StudentProfile.id).where(
                StudentProfile.enrollment_no == rec["enrollment_no"])).first():
            err("enrollment_no", "This enrollment number is already registered.")
            continue
        key = (dept.id, batch.id, rec["roll_number"])
        if key in seen_rolls:
            err("roll_number", "Duplicate roll number earlier in this file.")
            continue
        if rec["email"] in seen_emails:
            err("email", "Duplicate email earlier in this file.")
            continue
        if rec["phone"] in seen_phones:
            err("mobile", "Duplicate mobile earlier in this file.")
            continue
        try:
            check_unique(db, email=rec["email"], phone=rec["phone"], roll=rec["roll_number"],
                         department_id=dept.id, batch_id=batch.id)
        except HTTPException as exc:
            d = exc.detail if isinstance(exc.detail, dict) else {"field": "", "message": str(exc.detail)}
            err(d.get("field") or "row", d.get("message", "conflict"))
            continue
        seen_rolls.add(key)
        seen_emails.add(rec["email"])
        seen_phones.add(rec["phone"])
        seen_enrol.add(rec["enrollment_no"])
        valid.append({**rec, "dept": dept, "batch": batch, "section": sec, "year": year})

    created_ids: List[int] = []
    if not dry_run and valid:
        from app.routes.accounts import _username_for

        for rec in valid:
            u = User(username=_username_for(db, rec["email"]), display_name=rec["full_name"],
                     email=rec["email"], phone=rec["phone"], roll_number=rec["roll_number"],
                     role=Role.STUDENT, password_hash=hash_password(secrets.token_urlsafe(24)))
            db.add(u)
            db.flush()
            db.add(StudentProfile(
                user_id=u.id, section_id=rec["section"].id, department_id=rec["dept"].id,
                batch_id=rec["batch"].id, year=rec["year"], roll_number=rec["roll_number"],
                enrollment_no=rec["enrollment_no"], dob=rec["dob"], gender=rec["gender"],
                guardian_name=rec["guardian_name"], guardian_phone=rec["guardian_phone"],
                guardian_email=rec["guardian_email"], address=rec["address"],
                status=StudentStatus.APPROVED.value, decided_by_id=actor.id,
                decided_at=utcnow(), admission_year=rec["batch"].start_year))
            if send_invites:
                # Created accounts start with an unguessable password; the student sets their
                # own through a one-time link (valid 7 days).
                token = secrets.token_urlsafe(32)
                db.add(PasswordResetToken(user_id=u.id, token_hash=token_digest(token),
                                          expires_at=utcnow() + INVITE_TTL))
                notify(db, [u.id], "account_invite", "Your attendance portal account",
                       f"An account was created for you ({rec['roll_number']}). Use the link "
                       f"below to set your password; it is valid for 7 days. You can then sign "
                       f"in with your email or roll number.",
                       link=f"/reset-password?token={token}", email=True, sensitive=True)
            created_ids.append(u.id)
        audit(db, actor, "students_imported", target_type="import",
              detail={"source": source, "rows": len(rows), "created": len(created_ids),
                      "errors": len(errors)}, request=request)
        db.commit()
    return {"rows": len(rows), "valid": len(valid), "created": len(created_ids),
            "created_ids": created_ids, "dry_run": dry_run, "errors": errors}


@router.post("/students/import")
async def import_students(request: Request, file: UploadFile = File(...),
                          dry_run: bool = Form(False), send_invites: bool = Form(True),
                          db: Session = Depends(get_db), actor: User = Depends(require_staff)):
    data = await read_limited(file, 5 * MB)
    rows = _read_rows(data, file.filename or "")
    if not rows:
        raise HTTPException(422, "No student rows found under the header.")
    if len(rows) > 2000:
        raise HTTPException(413, "Import at most 2000 students per file.")
    return process_rows(db, actor, rows, dry_run=dry_run, send_invites=send_invites,
                        source=file.filename or "upload", request=request)


class NewStudentIn(BaseModel):
    full_name: str
    email: str
    mobile: str
    roll_number: str
    enrollment_no: str
    section_id: int
    dob: Optional[str] = None
    gender: Optional[str] = None
    guardian_name: Optional[str] = None
    guardian_phone: Optional[str] = None
    guardian_email: Optional[str] = None
    address: Optional[str] = None


@router.post("/students", status_code=201)
def create_student(body: NewStudentIn, request: Request, db: Session = Depends(get_db),
                   actor: User = Depends(require_staff)):
    """Add one student by hand: same validation and invite as the bulk import."""
    sec = assert_can_manage_section(db, actor, body.section_id)
    if sec.batch_id is None:
        raise HTTPException(422, "Assign this section to a batch first (Setup > Sections).")
    row = {"full_name": body.full_name, "email": body.email, "mobile": body.mobile,
           "roll_number": body.roll_number, "enrollment_no": body.enrollment_no,
           "branch_code": sec.program.department.code, "batch": sec.batch.label,
           "year": str(sec.year), "section": sec.name, "dob": body.dob or "",
           "gender": body.gender or "", "guardian_name": body.guardian_name or "",
           "guardian_phone": body.guardian_phone or "",
           "guardian_email": body.guardian_email or "", "address": body.address or ""}
    result = process_rows(db, actor, [row], dry_run=False, send_invites=True,
                          source="manual entry", request=request)
    if result["errors"]:
        e = result["errors"][0]
        raise HTTPException(409 if "already" in e["message"] else 422,
                            detail={"field": e["field"], "message": e["message"]})
    uid = result["created_ids"][0]
    user = db.get(User, uid)
    return profile_out(db, user, db.scalars(select(StudentProfile)
                                            .where(StudentProfile.user_id == uid)).first())
