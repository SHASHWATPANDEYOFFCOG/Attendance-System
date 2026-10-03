"""Phase 8: summary reports, exports and scoping, checked against known numbers."""
from __future__ import annotations

import csv
import io
from datetime import datetime, timedelta, timezone

import pytest

from test_portal_sessions import _db, klass  # noqa: F401


@pytest.fixture()
def history(api, klass):
    """Ten closed lectures for the class; student 0 attends 6 of 10 (+1 on leave, so 6/9)."""
    from app.models import AttendanceRecordRow, ClassSession

    db = _db()
    now = datetime.now(timezone.utc)
    s0, s1, s2, s3 = klass["students"]
    plan = {s0: ["present"] * 5 + ["late"] + ["absent"] * 3 + ["on_leave"],
            s1: ["present"] * 10, s2: ["absent"] * 10, s3: ["present"] * 8 + ["absent"] * 2}
    for i in range(10):
        s = ClassSession(title="CS301", room="R", owner_id=klass["teacher"], state="closed",
                         scheduled_start=now - timedelta(days=i + 1, hours=1),
                         scheduled_end=now - timedelta(days=i + 1),
                         subject_id=klass["subject"], section_id=klass["section"],
                         capture_method="manual")
        db.add(s)
        db.flush()
        for uid, statuses in plan.items():
            db.add(AttendanceRecordRow(session_id=s.id, user_id=uid, status=statuses[i],
                                       source="manual"))
    db.commit()
    db.close()
    return plan


def _row(rows, key, value):
    return next(r for r in rows if r[key] == value)


def test_student_wise_report_numbers(api, klass, history):
    s0 = klass["students"][0]
    r = api.get(f"/api/reports/student-wise/{s0}", headers=klass["teacher_h"]).json()
    overall = _row(r["rows"], "subject", "OVERALL")
    assert overall["lectures"] == 10 and overall["late"] == 1 and overall["on_leave"] == 1
    assert overall["pct"] == round(100 * 6 / 9, 1)                 # 66.7
    assert overall["needed"] == 3                                    # (6+3)/(9+3) = 75%


def test_defaulter_list_and_subject_wise(api, klass, history):
    s0, s1, s2, s3 = klass["students"]
    d = api.get("/api/reports/defaulter-list", headers=klass["teacher_h"]).json()
    ids = [r["student_id"] for r in d["rows"]]
    assert ids == [s2, s0]                                           # worst first; s3 is 80%
    assert _row(d["rows"], "student_id", s2)["pct"] == 0.0
    subj = api.get("/api/reports/subject-wise", headers=klass["teacher_h"]).json()["rows"]
    assert len(subj) == 1 and subj[0]["lectures"] == 10 and subj[0]["below"] == 2
    # present+late attended = 6 + 10 + 0 + 8 = 24 of counted 39 (one on_leave excluded)
    assert subj[0]["attended"] == 24 and subj[0]["pct"] == round(100 * 24 / 39, 1)


def test_daily_and_monthly_reports(api, klass, history):
    daily = api.get("/api/reports/daily", headers=klass["teacher_h"]).json()["rows"]
    assert sum(r["lectures"] for r in daily) == 10
    assert all(r["marked"] == 4 for r in daily)
    month = datetime.now(timezone.utc).strftime("%Y-%m")
    monthly = api.get(f"/api/reports/monthly?month={month}", headers=klass["teacher_h"]).json()
    assert monthly["rows"], "monthly report is empty"
    assert {r["student_id"] for r in monthly["rows"]} <= set(klass["students"])


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "pdf"])
def test_exports_open_and_match_the_screen(api, klass, history, fmt):
    screen = api.get("/api/reports/defaulter-list", headers=klass["teacher_h"]).json()
    r = api.get(f"/api/reports/defaulter-list?format={fmt}", headers=klass["teacher_h"])
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    if fmt == "csv":
        text = r.content.decode("utf-8-sig")
        assert text.startswith("# Attendance Register") or text.startswith("# ")
        body = [l for l in text.splitlines() if l and not l.startswith("#")]
        rows = list(csv.DictReader(io.StringIO("\n".join(body))))
        assert [x["Student"] for x in rows] == [x["student"] for x in screen["rows"]]
        assert [float(x["Attendance %"]) for x in rows] == [x["pct"] for x in screen["rows"]]
    elif fmt == "xlsx":
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(r.content)).active
        values = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        for x in screen["rows"]:
            assert x["student"] in values
        assert any("Defaulter list" == v for v in values)
    else:
        assert r.content.startswith(b"%PDF") and len(r.content) > 1000


def test_reports_are_scoped_to_the_teachers_classes(api, klass, history):
    for path in ("/api/reports/defaulter-list", "/api/reports/daily", "/api/reports/subject-wise",
                 "/api/analytics/classes", "/api/reports/attendance", "/api/reports/defaulters"):
        r = api.get(path, headers=klass["other_h"]).json()
        rows = r.get("rows", r.get("classes", r.get("defaulters")))
        assert rows == [], f"{path} leaked another teacher's class: {rows}"
    s0 = klass["students"][0]
    assert api.get(f"/api/reports/student-wise/{s0}", headers=klass["other_h"]).status_code == 403
    assert api.get("/api/reports/daily", headers=klass["h"][s0]).status_code == 403


def test_class_wise_dashboard_numbers(api, klass, history):
    r = api.get("/api/analytics/classes", headers=klass["teacher_h"]).json()
    assert len(r["classes"]) == 1
    assert r["classes"][0]["pct"] == round(100 * 24 / 39, 1)
    assert r["classes"][0]["lectures"] == 10
