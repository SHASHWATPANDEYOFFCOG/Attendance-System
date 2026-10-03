"""Filterable reports and their exports.

The properties worth guarding: filtering happens in SQL (so paging is real), an export always
matches the table it was generated from, department scope cannot be widened by a query
parameter, and every export carries provenance.
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime, timedelta, timezone

import pytest

from app.crypto import hash_password
from app.db import SessionLocal
from app.models import (
    AttendanceRecordRow,
    AuditLog,
    ClassSession,
    Department,
    Program,
    Role,
    Section,
    StudentProfile,
    Subject,
    User,
)
from app.reports_service import ReportFilters, to_csv, to_pdf, to_xlsx


@pytest.fixture()
def data(api):
    """Two departments, each with one subject, one section, two students and two sessions."""
    db = SessionLocal()
    pw = hash_password("password123")
    out: dict = {"students": {}, "sessions": {}}
    try:
        base = datetime.now(timezone.utc).replace(hour=9, minute=0, second=0, microsecond=0)
        for code in ("CSE", "ECE"):
            dept = Department(code=code, name=code)
            db.add(dept)
            db.commit()
            db.refresh(dept)
            prog = Program(department_id=dept.id, name=f"B.Tech {code}")
            db.add(prog)
            db.commit()
            db.refresh(prog)
            sec = Section(program_id=prog.id, name="A", year=3, semester=5)
            subj = Subject(code=f"{code}301", name=f"{code} Core", department_id=dept.id)
            prof = User(username=f"prof.{code}", display_name=f"Prof {code}",
                        role=Role.FACULTY, password_hash=pw, department_id=dept.id)
            db.add_all([sec, subj, prof])
            db.commit()
            for o in (sec, subj, prof):
                db.refresh(o)

            uids = []
            for i in range(2):
                u = User(username=f"{code.lower()}{i}", display_name=f"{code} Student {i}",
                         role=Role.STUDENT, password_hash=pw, roll_number=f"{code}{i:03d}")
                db.add(u)
                db.commit()
                db.refresh(u)
                db.add(StudentProfile(user_id=u.id, section_id=sec.id, admission_year=2021))
                uids.append(u.id)
            db.commit()

            # Two days: yesterday and three days ago.
            sess_ids = []
            for offset in (1, 3):
                s = ClassSession(
                    title=f"{code} Lecture", room=f"{code}-1", owner_id=prof.id,
                    scheduled_start=base - timedelta(days=offset),
                    scheduled_end=base - timedelta(days=offset) + timedelta(minutes=50),
                    mode="kiosk", grace_minutes=10, min_presence_seconds=1500,
                    subject_id=subj.id, section_id=sec.id, capture_method="manual",
                )
                db.add(s)
                db.commit()
                db.refresh(s)
                sess_ids.append(s.id)
                for j, uid in enumerate(uids):
                    db.add(AttendanceRecordRow(
                        session_id=s.id, user_id=uid,
                        status="present" if j == 0 else "absent", source="manual"))
                db.commit()

            out["students"][code] = uids
            out["sessions"][code] = sess_ids
            out[f"{code}_dept"] = dept.id
            out[f"{code}_section"] = sec.id
            out[f"{code}_subject"] = subj.id
            out[f"{code}_prof"] = prof.id
        return out
    finally:
        db.close()


@pytest.fixture()
def staff(api, make_user):
    _uid, headers = make_user("root", role="super_admin")
    return headers


# ------------------------------------------------------------------ filtering

def test_report_returns_rows_and_pagination_meta(api, staff, data):
    body = api.get("/api/reports/attendance", headers=staff).json()
    assert body["meta"]["total"] == 8            # 2 depts x 2 sessions x 2 students
    assert len(body["rows"]) == 8
    assert body["meta"]["pages"] == 1


def test_paging_is_done_in_sql(api, staff, data):
    p1 = api.get("/api/reports/attendance?page=1&page_size=3", headers=staff).json()
    p2 = api.get("/api/reports/attendance?page=2&page_size=3", headers=staff).json()
    assert len(p1["rows"]) == 3 and len(p2["rows"]) == 3
    assert p1["meta"]["pages"] == 3
    ids = {r["record_id"] for r in p1["rows"]} & {r["record_id"] for r in p2["rows"]}
    assert not ids, "pages overlap — offset is not being applied"


def test_page_size_is_capped(api, staff, data):
    body = api.get("/api/reports/attendance?page_size=99999", headers=staff).json()
    assert body["meta"]["page_size"] <= 200


def test_filter_by_status(api, staff, data):
    body = api.get("/api/reports/attendance?status=absent", headers=staff).json()
    assert body["meta"]["total"] == 4
    assert all(r["status"] == "absent" for r in body["rows"])


def test_filter_by_subject_and_section(api, staff, data):
    body = api.get(f"/api/reports/attendance?subject_id={data['CSE_subject']}",
                   headers=staff).json()
    assert body["meta"]["total"] == 4
    assert all(r["subject_code"] == "CSE301" for r in body["rows"])


def test_filter_by_student(api, staff, data):
    uid = data["students"]["CSE"][0]
    body = api.get(f"/api/reports/attendance?student_id={uid}", headers=staff).json()
    assert body["meta"]["total"] == 2


def test_filter_by_date_range_is_inclusive(api, staff, data):
    yesterday = (datetime.now(timezone.utc) - timedelta(days=1)).date()
    body = api.get(f"/api/reports/attendance?date_from={yesterday}&date_to={yesterday}",
                   headers=staff).json()
    assert body["meta"]["total"] == 4, "the end date must be inclusive"
    assert all(r["date"] == yesterday.isoformat() for r in body["rows"])


def test_search_matches_name_or_roll(api, staff, data):
    body = api.get("/api/reports/attendance?search=CSE0", headers=staff).json()
    assert body["meta"]["total"] > 0
    assert all("CSE" in (r["roll_number"] or "") for r in body["rows"])


def test_a_malformed_date_is_rejected_clearly(api, staff, data):
    r = api.get("/api/reports/attendance?date_from=last-tuesday", headers=staff)
    assert r.status_code == 422
    assert "expected YYYY-MM-DD" in r.text


# ------------------------------------------------------------------- scoping

def test_hod_sees_only_their_department(api, make_user, data):
    _uid, hod = make_user("hod.cse", role="admin", department_id=data["CSE_dept"])
    body = api.get("/api/reports/attendance", headers=hod).json()
    assert body["meta"]["total"] == 4
    assert all(r["subject_code"] == "CSE301" for r in body["rows"])


def test_hod_cannot_widen_scope_with_a_query_parameter(api, make_user, data):
    _uid, hod = make_user("hod.cse", role="admin", department_id=data["CSE_dept"])
    assert api.get(f"/api/reports/attendance?department_id={data['ECE_dept']}",
                   headers=hod).status_code == 403


def test_reports_are_staff_only(api, make_user, data):
    _uid, student = make_user("pupil")
    for path in ("/api/reports/attendance", "/api/reports/daily-absentees",
                 "/api/reports/export"):
        assert api.get(path, headers=student).status_code == 403, path


# ---------------------------------------------------------------- absentees

def test_daily_absentees_groups_by_student(api, staff, data):
    yesterday = (datetime.now(timezone.utc) - timedelta(days=1)).date()
    body = api.get(f"/api/reports/daily-absentees?date={yesterday}", headers=staff).json()
    assert body["count"] == 2                    # one absent student per department
    assert all(e["periods_missed"] >= 1 for e in body["absentees"])
    assert all(e["subjects"] for e in body["absentees"])


def test_daily_absentees_on_a_quiet_day_is_empty(api, staff, data):
    body = api.get("/api/reports/daily-absentees?date=2020-01-01", headers=staff).json()
    assert body["count"] == 0 and body["absentees"] == []


# ------------------------------------------------------------------ exports

def test_csv_export_matches_the_table(api, staff, data):
    table = api.get("/api/reports/attendance?status=absent&page_size=200",
                    headers=staff).json()
    r = api.get("/api/reports/export?format=csv&status=absent", headers=staff)
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]

    lines = [l for l in r.text.splitlines() if l and not l.startswith("#")]
    reader = csv.DictReader(io.StringIO("\n".join(lines)))
    rows = list(reader)
    assert len(rows) == table["meta"]["total"]
    assert all(row["Status"] == "absent" for row in rows)


def test_exports_carry_provenance(api, staff, data):
    text = api.get("/api/reports/export?format=csv", headers=staff).text
    assert "# Attendance report" in text
    assert "# Generated" in text
    assert "row(s)" in text
    assert "Filters:" in text


def test_export_records_the_filters_it_used(api, staff, data):
    text = api.get("/api/reports/export?format=csv&status=late", headers=staff).text
    assert "Status=late" in text


def test_xlsx_export_is_a_real_workbook(api, staff, data):
    from openpyxl import load_workbook

    r = api.get("/api/reports/export?format=xlsx", headers=staff)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.title == "Attendance"
    flat = [str(c.value) for row in ws.iter_rows(max_row=8) for c in row if c.value]
    assert any("Attendance report" in v for v in flat)
    assert any(v == "Roll" for v in flat), "header row missing"


def test_pdf_export_is_a_real_pdf(api, staff, data):
    r = api.get("/api/reports/export?format=pdf", headers=staff)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:5] == b"%PDF-", "not a PDF"
    assert len(r.content) > 1000


def test_times_are_rendered_in_the_org_timezone_not_utc(api, staff, data):
    """Stored UTC, rendered local. A 09:00 IST lecture shown as 03:30 is not actionable."""
    from app.reports_service import ReportFilters, query_records

    db = SessionLocal()
    try:
        utc_rows = query_records(db, ReportFilters(), None, limit=5, tz="UTC")
        ist_rows = query_records(db, ReportFilters(), None, limit=5, tz="Asia/Kolkata")
        assert utc_rows and ist_rows

        u, i = utc_rows[0], ist_rows[0]
        assert u["record_id"] == i["record_id"]
        uh, um = (int(x) for x in u["time"].split(":"))
        ih, im = (int(x) for x in i["time"].split(":"))
        offset = ((ih * 60 + im) - (uh * 60 + um)) % (24 * 60)
        assert offset == 330, f"expected +05:30, got {offset} minutes"
    finally:
        db.close()


def test_an_unknown_timezone_falls_back_to_utc(api, staff, data):
    """A bad config value must degrade to UTC, not blank the whole report."""
    from app.reports_service import ReportFilters, query_records

    db = SessionLocal()
    try:
        rows = query_records(db, ReportFilters(), None, limit=1, tz="Mars/Olympus_Mons")
        assert rows and rows[0]["time"]
    finally:
        db.close()


def test_pdf_has_a_lower_row_limit_than_the_others(api, staff, data):
    """xhtml2pdf's cost is superlinear (25s for 2 000 rows on this machine), so a large PDF
    must be refused up front rather than hanging the request until it times out."""
    from app.reports_service import EXPORT_ROW_LIMITS, MAX_EXPORT_ROWS

    assert EXPORT_ROW_LIMITS["pdf"] < EXPORT_ROW_LIMITS["csv"]
    assert EXPORT_ROW_LIMITS["csv"] == MAX_EXPORT_ROWS


def test_an_oversized_export_is_refused_with_a_useful_message(api, staff, data, monkeypatch):
    import app.reports_service as rs

    monkeypatch.setitem(rs.EXPORT_ROW_LIMITS, "pdf", 2)
    r = api.get("/api/reports/export?format=pdf", headers=staff)
    assert r.status_code == 413
    body = r.json()["detail"]
    assert body["limit"] == 2 and body["total"] == 8
    assert "XLSX/CSV" in body["message"], "the message must name a way forward"

    # The other formats are unaffected by the PDF ceiling.
    assert api.get("/api/reports/export?format=csv", headers=staff).status_code == 200


def test_an_unknown_export_format_is_refused(api, staff, data):
    r = api.get("/api/reports/export?format=docx", headers=staff)
    assert r.status_code == 422
    assert "csv, xlsx or pdf" in r.text


def test_export_respects_department_scope(api, make_user, data):
    _uid, hod = make_user("hod.cse", role="admin", department_id=data["CSE_dept"])
    text = api.get("/api/reports/export?format=csv", headers=hod).text
    assert "CSE301" in text
    assert "ECE301" not in text, "an HOD exported another department's rows"


def test_export_is_audit_logged(api, staff, data):
    api.get("/api/reports/export?format=csv", headers=staff)
    db = SessionLocal()
    try:
        entry = db.query(AuditLog).filter(AuditLog.action == "export_report").one()
        assert "format=csv" in entry.detail
    finally:
        db.close()


def test_exporters_handle_an_empty_result(api, staff, data):
    """An empty report is a normal answer, not an error, and must still be openable."""
    filters = ReportFilters(status="left_early")
    assert "Attendance report" in to_csv([], "Demo College", filters, 0)
    assert to_xlsx([], "Demo College", filters, 0)[:2] == b"PK"     # zip magic
    assert to_pdf([], "Demo College", filters, 0)[:5] == b"%PDF-"


def test_pdf_escapes_html_in_names(api, staff, data):
    """A student named with angle brackets must not break or inject into the document."""
    rows = [{"date": "2026-01-01", "time": "09:00", "roll_number": "R1",
             "student": "<script>alert(1)</script>", "subject_code": "CS1",
             "section": "A", "status": "present", "source": "manual",
             "confidence": 0.9, "overridden": False}]
    pdf = to_pdf(rows, "Demo", ReportFilters(), 1)
    assert pdf[:5] == b"%PDF-"
