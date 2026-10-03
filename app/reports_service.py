"""Filterable, paginated attendance reports and their exports.

Two rules shape this module:

* **Filtering and paging happen in SQL.** A term's register is ~22,000 rows for a small college
  and grows linearly; materialising it to filter in Python would work in the demo and fall over
  in production. Every filter is a WHERE clause and every page is a LIMIT/OFFSET.
* **Exports stream from the same query as the screen.** The CSV, the workbook and the PDF are
  built from `query_records`, so an export can never disagree with the table the user was
  looking at when they clicked it.

Exports carry the organisation name and a generated-at stamp, because a printed attendance
report with no provenance is not evidence of anything.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from typing import List, Optional, Sequence, Tuple

from sqlalchemy import Select, func, or_, select
from sqlalchemy.orm import Session

from app.models import (
    COUNTS_PRESENT,
    EXCLUDED_FROM_TOTAL,
    AttendanceRecordRow,
    ClassSession,
    Department,
    Program,
    Section,
    StudentProfile,
    Subject,
    User,
    utcnow,
)
from occlubio.utils import get_logger

log = get_logger("app.reports")

MAX_PAGE_SIZE = 200
#: Hard ceiling on an export. Past this a report is a data dump, not a report, and should be a
#: scheduled extract instead of a synchronous HTTP response.
MAX_EXPORT_ROWS = 20_000

#: Per-format ceilings. PDF gets its own, far lower one because xhtml2pdf's cost is
#: SUPERLINEAR in row count -- measured on this machine:
#:     100 rows  3.3s  |  1 000 rows  8.9s  |  2 000 rows  24.9s
#: versus XLSX at 0.34s and CSV at 0.004s for the same 2 000. Allowing a 20 000-row PDF would
#: hang the worker for minutes and then time out, which is worse than refusing up front.
#: A PDF is a printable artefact; nobody prints twenty thousand rows.
EXPORT_ROW_LIMITS = {"csv": MAX_EXPORT_ROWS, "xlsx": MAX_EXPORT_ROWS, "pdf": 1_000}


@dataclass
class ReportFilters:
    """Everything the filter builder can set. All optional; all applied in SQL."""

    date_from: Optional[date] = None
    date_to: Optional[date] = None
    department_id: Optional[int] = None
    section_id: Optional[int] = None
    subject_id: Optional[int] = None
    faculty_id: Optional[int] = None
    student_id: Optional[int] = None
    status: Optional[str] = None
    search: Optional[str] = None

    def describe(self) -> List[Tuple[str, str]]:
        """Human-readable filter list for an export header."""
        out: List[Tuple[str, str]] = []
        if self.date_from:
            out.append(("From", self.date_from.isoformat()))
        if self.date_to:
            out.append(("To", self.date_to.isoformat()))
        for label, value in (
            ("Department", self.department_id), ("Section", self.section_id),
            ("Subject", self.subject_id), ("Faculty", self.faculty_id),
            ("Student", self.student_id),
        ):
            if value is not None:
                out.append((label, f"#{value}"))
        if self.status:
            out.append(("Status", self.status))
        if self.search:
            out.append(("Search", self.search))
        return out


def _effective():
    return func.coalesce(AttendanceRecordRow.overridden_status, AttendanceRecordRow.status)


def _base_query(filters: ReportFilters, department_scope: Optional[int]) -> Select:
    stmt = (
        select(
            AttendanceRecordRow.id,
            ClassSession.scheduled_start,
            User.display_name,
            User.roll_number,
            Subject.code,
            Subject.name,
            Section.name,
            _effective().label("status"),
            AttendanceRecordRow.status.label("derived_status"),
            AttendanceRecordRow.overridden_status,
            AttendanceRecordRow.avg_confidence,
            AttendanceRecordRow.source,
            AttendanceRecordRow.review_reason,
        )
        .join(ClassSession, ClassSession.id == AttendanceRecordRow.session_id)
        .join(User, User.id == AttendanceRecordRow.user_id)
        .outerjoin(Subject, Subject.id == ClassSession.subject_id)
        .outerjoin(Section, Section.id == ClassSession.section_id)
        .where(AttendanceRecordRow.deleted_at.is_(None))
        .where(ClassSession.deleted_at.is_(None))
        .where(User.deleted_at.is_(None))
    )

    # The caller's department scope is applied FIRST and cannot be widened by a filter: an HOD
    # passing ?department_id= for someone else is refused at the route, and this is the backstop.
    if department_scope is not None:
        stmt = stmt.where(Subject.department_id == department_scope)
    elif filters.department_id is not None:
        stmt = stmt.where(Subject.department_id == filters.department_id)

    if filters.date_from:
        stmt = stmt.where(ClassSession.scheduled_start >= _utc(filters.date_from))
    if filters.date_to:
        stmt = stmt.where(ClassSession.scheduled_start < _utc(filters.date_to) + timedelta(days=1))
    if filters.section_id is not None:
        stmt = stmt.where(ClassSession.section_id == filters.section_id)
    if filters.subject_id is not None:
        stmt = stmt.where(ClassSession.subject_id == filters.subject_id)
    if filters.faculty_id is not None:
        stmt = stmt.where(ClassSession.owner_id == filters.faculty_id)
    if filters.student_id is not None:
        stmt = stmt.where(AttendanceRecordRow.user_id == filters.student_id)
    if filters.status:
        stmt = stmt.where(_effective() == filters.status)
    if filters.search:
        like = f"%{filters.search}%"
        stmt = stmt.where(or_(User.display_name.ilike(like), User.roll_number.ilike(like)))
    return stmt


def _utc(d: date) -> datetime:
    return datetime.combine(d, datetime.min.time()).replace(tzinfo=timezone.utc)


def count_records(db: Session, filters: ReportFilters,
                  department_scope: Optional[int]) -> int:
    inner = _base_query(filters, department_scope).subquery()
    return int(db.scalar(select(func.count()).select_from(inner)) or 0)


def _zone(tz: str):
    """Resolve a timezone name, falling back to UTC rather than failing a whole report."""
    try:
        from zoneinfo import ZoneInfo

        return ZoneInfo(tz)
    except Exception:                                   # noqa: BLE001 - bad name or no tzdata
        log.warning("unknown timezone %r; rendering report times in UTC", tz)
        return timezone.utc


def query_records(db: Session, filters: ReportFilters, department_scope: Optional[int],
                  *, page: int = 1, page_size: int = 50,
                  limit: Optional[int] = None, tz: str = "UTC") -> List[dict]:
    """One page of the report. `limit` overrides paging for exports.

    Timestamps are stored UTC and rendered in the organisation's zone here — a report that
    shows a 09:00 lecture as 03:30 is not a report anyone can act on. The DATE is converted
    too, so a late-evening class does not appear on the wrong day.
    """
    zone = _zone(tz)
    stmt = _base_query(filters, department_scope).order_by(
        ClassSession.scheduled_start.desc(), User.roll_number)

    if limit is not None:
        stmt = stmt.limit(limit)
    else:
        page = max(1, page)
        page_size = max(1, min(MAX_PAGE_SIZE, page_size))
        stmt = stmt.limit(page_size).offset((page - 1) * page_size)

    rows = []
    for r in db.execute(stmt).all():
        local = r[1].astimezone(zone) if r[1] else None
        rows.append({
            "record_id": r[0],
            "date": local.date().isoformat() if local else None,
            "time": local.strftime("%H:%M") if local else None,
            "student": r[2],
            "roll_number": r[3],
            "subject_code": r[4],
            "subject_name": r[5],
            "section": r[6],
            "status": r[7],
            "derived_status": r[8],
            "overridden": bool(r[9]),
            "confidence": round(r[10], 3) if r[10] else None,
            "source": r[11],
            "review_reason": r[12],
        })
    return rows


def daily_absentees(db: Session, on: date, department_scope: Optional[int],
                    *, section_id: Optional[int] = None, tz: str = "UTC") -> List[dict]:
    """Everyone marked absent on one day, grouped per student.

    The list a front office actually uses: one row per student with the periods they missed,
    not one row per missed period.
    """
    filters = ReportFilters(date_from=on, date_to=on, status="absent", section_id=section_id)
    rows = query_records(db, filters, department_scope, limit=MAX_EXPORT_ROWS, tz=tz)

    per_student: dict = {}
    for r in rows:
        key = r["roll_number"] or r["student"]
        entry = per_student.setdefault(key, {
            "student": r["student"], "roll_number": r["roll_number"],
            "section": r["section"], "periods_missed": 0, "subjects": [],
        })
        entry["periods_missed"] += 1
        if r["subject_code"] and r["subject_code"] not in entry["subjects"]:
            entry["subjects"].append(r["subject_code"])

    out = sorted(per_student.values(), key=lambda e: -e["periods_missed"])
    return out


# ==========================================================================================
# EXPORTS
# ==========================================================================================

COLUMNS = [
    ("date", "Date"),
    ("time", "Time"),
    ("roll_number", "Roll"),
    ("student", "Student"),
    ("subject_code", "Subject"),
    ("section", "Section"),
    ("status", "Status"),
    ("source", "Source"),
    ("confidence", "Confidence"),
    ("overridden", "Overridden"),
]


def _header_lines(org: str, filters: ReportFilters, total: int, tz: str) -> List[str]:
    stamp = utcnow().strftime("%Y-%m-%d %H:%M UTC")
    lines = [org, "Attendance report", f"Generated {stamp} (rendered in {tz})",
             f"{total} row(s)"]
    described = filters.describe()
    if described:
        lines.append("Filters: " + " · ".join(f"{k}={v}" for k, v in described))
    else:
        lines.append("Filters: none (all records)")
    return lines


def to_csv(rows: Sequence[dict], org: str, filters: ReportFilters, total: int,
           tz: str = "UTC") -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    # Provenance first, as comment rows: a CSV opened six months later must still say what it is.
    for line in _header_lines(org, filters, total, tz):
        writer.writerow([f"# {line}"])
    writer.writerow([])
    writer.writerow([label for _key, label in COLUMNS])
    for r in rows:
        writer.writerow([r.get(key, "") for key, _label in COLUMNS])
    return buf.getvalue()


def to_xlsx(rows: Sequence[dict], org: str, filters: ReportFilters, total: int,
            tz: str = "UTC") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    for i, line in enumerate(_header_lines(org, filters, total, tz), start=1):
        cell = ws.cell(row=i, column=1, value=line)
        cell.font = Font(bold=(i == 1), size=13 if i == 1 else 10,
                         color="16304F" if i == 1 else "6E7468")
    head_row = len(_header_lines(org, filters, total, tz)) + 2

    for c, (_key, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=head_row, column=c, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16304F")
        cell.alignment = Alignment(horizontal="left")

    status_fill = {
        "present": "E4F0E9", "late": "FBF0DC", "absent": "F6E2DE", "excused": "EDEEEB",
    }
    for r_i, row in enumerate(rows, start=head_row + 1):
        for c_i, (key, _label) in enumerate(COLUMNS, start=1):
            value = row.get(key)
            if isinstance(value, bool):
                value = "yes" if value else ""
            cell = ws.cell(row=r_i, column=c_i, value=value)
            if key == "status" and value in status_fill:
                cell.fill = PatternFill("solid", fgColor=status_fill[value])

    for c, (key, label) in enumerate(COLUMNS, start=1):
        width = max(len(label) + 2,
                    *(len(str(r.get(key, ""))) + 2 for r in rows[:200])) if rows else len(label) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(32, width)
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(rows: Sequence[dict], org: str, filters: ReportFilters, total: int,
           tz: str = "UTC") -> bytes:
    """Render via xhtml2pdf: pure Python, no headless browser and no system binaries."""
    from xhtml2pdf import pisa

    def esc(v) -> str:
        return (str("" if v is None else v)
                .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

    header = "".join(f"<div class='sub'>{esc(l)}</div>" for l in _header_lines(org, filters, total, tz)[1:])
    body_rows = "".join(
        "<tr>" + "".join(
            f"<td class='{ 'st-' + str(r.get('status')) if key == 'status' else '' }'>"
            f"{esc('yes' if r.get(key) is True else ('' if r.get(key) is False else r.get(key)))}</td>"
            for key, _label in COLUMNS
        ) + "</tr>"
        for r in rows
    )

    html = f"""
    <html><head><style>
      @page {{ size: A4 landscape; margin: 1.4cm; }}
      body {{ font-family: Helvetica; font-size: 8pt; color: #16304F; }}
      h1 {{ font-size: 15pt; margin: 0 0 2pt; letter-spacing: 1pt; }}
      .sub {{ font-size: 7.5pt; color: #6E7468; }}
      hr {{ border: 0; border-top: 1.5px solid #16304F; margin: 6pt 0; }}
      table {{ width: 100%; border-collapse: collapse; }}
      th {{ background: #16304F; color: #fff; text-align: left; padding: 3pt 4pt; font-size: 7.5pt; }}
      td {{ border-bottom: 0.5px solid #D9DDD3; padding: 2.5pt 4pt; }}
      .st-present {{ color: #2E6B4F; }} .st-late {{ color: #A8731A; }}
      .st-absent {{ color: #9A3B2E; }} .st-excused {{ color: #6E7468; }}
    </style></head><body>
      <h1>{esc(org)}</h1>{header}
      <hr/>
      <table>
        <thead><tr>{''.join(f'<th>{esc(l)}</th>' for _k, l in COLUMNS)}</tr></thead>
        <tbody>{body_rows or "<tr><td colspan='10'>No rows matched these filters.</td></tr>"}</tbody>
      </table>
    </body></html>
    """

    out = io.BytesIO()
    result = pisa.CreatePDF(io.StringIO(html), dest=out)
    if result.err:
        raise RuntimeError("PDF generation failed")
    return out.getvalue()


EXPORTERS = {
    "csv": (to_csv, "text/csv", "csv"),
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": (to_pdf, "application/pdf", "pdf"),
}
