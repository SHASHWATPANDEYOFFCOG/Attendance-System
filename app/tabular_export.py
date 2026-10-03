"""Generic table export: CSV, Excel and PDF from one (columns, rows) description.

Every export carries the college name, the report title, when it was generated and what it was
filtered by, so a printout found in a drawer months later still says what it is. Built from the
same rows the screen shows, so the two can never disagree.
"""
from __future__ import annotations

import csv
import io
from typing import Any, List, Sequence, Tuple

from app.models import utcnow

Column = Tuple[str, str]          # (key, label)

STATUS_COLORS = {
    "present": ("2E6B4F", "E4F0E9"), "late": ("A8731A", "FBF0DC"), "absent": ("9A3B2E", "F6E2DE"),
    "on_leave": ("3D6DA8", "E3ECF7"), "excused": ("6E7468", "EDEEEB"), "on_duty": ("2E6B4F", "E4F0E9"),
}


def _cell(v: Any) -> Any:
    if isinstance(v, bool):
        return "yes" if v else ""
    return "" if v is None else v


def header_lines(org: str, title: str, meta: Sequence[str], tz: str) -> List[str]:
    return [org, title, f"Generated {utcnow().strftime('%Y-%m-%d %H:%M')} UTC "
                        f"(times shown in {tz})", *meta]


def to_csv(org: str, title: str, meta: Sequence[str], columns: Sequence[Column],
           rows: Sequence[dict], tz: str) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    for line in header_lines(org, title, meta, tz):
        w.writerow([f"# {line}"])
    w.writerow([])
    w.writerow([label for _k, label in columns])
    for r in rows:
        w.writerow([_cell(r.get(k)) for k, _l in columns])
    # UTF-8 with BOM so Excel on Windows opens names with accents correctly.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(org: str, title: str, meta: Sequence[str], columns: Sequence[Column],
            rows: Sequence[dict], tz: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"
    lines = header_lines(org, title, meta, tz)
    for i, line in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=line)
        c.font = Font(bold=i <= 2, size=14 if i == 1 else 10, color="16304F" if i <= 2 else "6E7468")
    head = len(lines) + 2
    for ci, (_k, label) in enumerate(columns, start=1):
        c = ws.cell(row=head, column=ci, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="16304F")
        c.alignment = Alignment(horizontal="left")
    for ri, r in enumerate(rows, start=head + 1):
        for ci, (k, _l) in enumerate(columns, start=1):
            v = _cell(r.get(k))
            c = ws.cell(row=ri, column=ci, value=v)
            if k == "status" and v in STATUS_COLORS:
                c.fill = PatternFill("solid", fgColor=STATUS_COLORS[v][1])
    for ci, (k, label) in enumerate(columns, start=1):
        width = max([len(label)] + [len(str(_cell(r.get(k)))) for r in rows[:300]]) + 2
        ws.column_dimensions[get_column_letter(ci)].width = min(40, width)
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _esc(v: Any) -> str:
    return (str(_cell(v)).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def to_pdf(org: str, title: str, meta: Sequence[str], columns: Sequence[Column],
           rows: Sequence[dict], tz: str, landscape: bool = True) -> bytes:
    from xhtml2pdf import pisa

    lines = header_lines(org, title, meta, tz)
    status_css = " ".join(f".st-{k} {{ color: #{fg}; }}" for k, (fg, _bg) in STATUS_COLORS.items())
    body = "".join(
        "<tr>" + "".join(
            f"<td class='{('st-' + str(r.get(k))) if k == 'status' else ''}'>{_esc(r.get(k))}</td>"
            for k, _l in columns) + "</tr>"
        for r in rows)
    html = f"""
    <html><head><meta charset="utf-8"/><style>
      @page {{ size: A4 {'landscape' if landscape else 'portrait'}; margin: 1.3cm;
               @frame footer {{ -pdf-frame-content: foot; bottom: 0.6cm; height: 0.6cm;
                                left: 1.3cm; right: 1.3cm; }} }}
      body {{ font-family: Helvetica; font-size: 8.5pt; color: #16304F; }}
      h1 {{ font-size: 16pt; margin: 0; letter-spacing: 0.5pt; }}
      h2 {{ font-size: 11pt; margin: 2pt 0 4pt; color: #9A3B2E; }}
      .sub {{ font-size: 7.5pt; color: #6E7468; }}
      hr {{ border: 0; border-top: 1.5px solid #16304F; margin: 6pt 0; }}
      table {{ width: 100%; border-collapse: collapse; }}
      th {{ background: #16304F; color: #fff; text-align: left; padding: 3pt 4pt; font-size: 7.5pt; }}
      td {{ border-bottom: 0.5px solid #D9DDD3; padding: 2.5pt 4pt; }}
      {status_css}
      #foot {{ font-size: 7pt; color: #6E7468; text-align: right; }}
    </style></head><body>
      <h1>{_esc(lines[0])}</h1><h2>{_esc(lines[1])}</h2>
      {''.join(f"<div class='sub'>{_esc(l)}</div>" for l in lines[2:])}
      <hr/>
      <table><thead><tr>{''.join(f'<th>{_esc(l)}</th>' for _k, l in columns)}</tr></thead>
      <tbody>{body or f"<tr><td colspan='{len(columns)}'>No rows.</td></tr>"}</tbody></table>
      <div id="foot">{_esc(lines[0])} &middot; page <pdf:pagenumber/> of <pdf:pagecount/></div>
    </body></html>"""
    out = io.BytesIO()
    result = pisa.CreatePDF(io.StringIO(html), dest=out, encoding="utf-8")
    if result.err:
        raise RuntimeError("PDF generation failed")
    return out.getvalue()


FORMATS = {
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": ("application/pdf", "pdf"),
}


def export(fmt: str, *, org: str, title: str, meta: Sequence[str], columns: Sequence[Column],
           rows: Sequence[dict], tz: str) -> Tuple[bytes, str, str]:
    fmt = fmt.lower()
    if fmt not in FORMATS:
        raise ValueError(f"unknown format {fmt!r}; expected csv, xlsx or pdf")
    fn = {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf}[fmt]
    media, ext = FORMATS[fmt]
    return fn(org, title, meta, columns, rows, tz), media, ext
